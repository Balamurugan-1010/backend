import os
import uuid
from datetime import datetime
from typing import List, Dict, Any, Tuple
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# On Vercel serverless, only /tmp is writable
if os.getenv("VERCEL"):
    DOCUMENTS_DIR = "/tmp/generated_documents"
else:
    DOCUMENTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "generated_documents"))
os.makedirs(DOCUMENTS_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex: str):
    """Set table cell background color."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_shortlist_docx(
    company_name: str,
    college_name: str,
    role_name: str,
    candidates_data: List[Dict[str, Any]],
    include_full_details: bool = True
) -> Tuple[str, str]:
    """
    Generates a beautifully formatted Word document (.docx) with shortlisted candidates.
    """
    doc = Document()

    # Set standard margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("TALENTMATCH")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(37, 99, 235)  # #2563EB Blue

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("AI-POWERED CANDIDATE SHORTLIST REPORT")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)  # Slate secondary

    # 2. Drive & Job Role Metadata Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        ("Company / Recruiter:", company_name),
        ("College / Institute:", college_name),
        ("Job Role Target:", role_name),
        ("Generated Date:", datetime.utcnow().strftime("%B %d, %Y"))
    ]

    for row_idx, (label, val) in enumerate(meta_items):
        row = meta_table.rows[row_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.8)

        set_cell_background(cell_lbl, "EFF6FF")
        set_cell_background(cell_val, "F8FAFC")

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(2)
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = RGBColor(30, 41, 59)

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(2)
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(10)
        r_val.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3. Candidate Summary Table Heading
    sec_heading = doc.add_paragraph()
    sec_heading.paragraph_format.space_before = Pt(8)
    sec_heading.paragraph_format.space_after = Pt(6)
    r_h = sec_heading.add_run("Candidate Summary Table")
    r_h.font.name = "Arial"
    r_h.font.size = Pt(14)
    r_h.font.bold = True
    r_h.font.color.rgb = RGBColor(29, 78, 216)

    # Summary Table
    headers = ["Rank", "Candidate Name", "Department", "CGPA", "Match Score", "Recruitment Status"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    widths = [Inches(0.6), Inches(2.0), Inches(1.8), Inches(0.7), Inches(1.1), Inches(1.1)]

    for idx, text in enumerate(headers):
        hdr_cells[idx].width = widths[idx]
        set_cell_background(hdr_cells[idx], "2563EB")
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for rank, cand in enumerate(candidates_data, start=1):
        row_cells = table.add_row().cells
        bg_color = "FFFFFF" if rank % 2 != 0 else "F8FAFC"

        vals = [
            f"#{rank}",
            f"{cand.get('name', '')} ({cand.get('candidate_code', '')})",
            cand.get("department", ""),
            f"{float(cand.get('cgpa', 0.0)):.2f}",
            f"{int(round(float(cand.get('match_score', 0.0))))}%",
            cand.get("status", "Shortlisted")
        ]

        for c_idx, val in enumerate(vals):
            row_cells[c_idx].width = widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            if c_idx == 4:
                r.font.bold = True
                r.font.color.rgb = RGBColor(37, 99, 235)

    # 4. Detailed Profiles Section
    if include_full_details:
        doc.add_page_break()

        det_heading = doc.add_paragraph()
        det_heading.paragraph_format.space_before = Pt(10)
        det_heading.paragraph_format.space_after = Pt(8)
        r_det = det_heading.add_run("Detailed Candidate Profiles")
        r_det.font.name = "Arial"
        r_det.font.size = Pt(16)
        r_det.font.bold = True
        r_det.font.color.rgb = RGBColor(29, 78, 216)

        for rank, cand in enumerate(candidates_data, start=1):
            cand_p = doc.add_paragraph()
            cand_p.paragraph_format.space_before = Pt(10)
            cand_p.paragraph_format.space_after = Pt(2)
            rc = cand_p.add_run(f"Rank {rank}: {cand.get('name')}  |  ID: {cand.get('candidate_code')}")
            rc.font.bold = True
            rc.font.size = Pt(12)
            rc.font.color.rgb = RGBColor(37, 99, 235)

            # Contact & Academic summary
            info_p = doc.add_paragraph()
            info_p.paragraph_format.space_after = Pt(3)
            info_p.add_run(f"Email: {cand.get('email')}  •  Phone: {cand.get('phone')}  •  CGPA: {cand.get('cgpa')}  •  Dept: {cand.get('department')}\n")
            score_run = info_p.add_run(f"AI Match Score: {int(round(float(cand.get('match_score', 0))))}%  •  Status: {cand.get('status', 'Shortlisted')}")
            score_run.font.bold = True
            score_run.font.color.rgb = RGBColor(16, 185, 129)  # Success green

            # Matched Skills & Gap
            matched = cand.get("matched_skills", [])
            missing = cand.get("missing_skills", [])

            skill_p = doc.add_paragraph()
            skill_p.paragraph_format.space_after = Pt(2)
            r_ms = skill_p.add_run("Matched Skills: ")
            r_ms.font.bold = True
            skill_p.add_run(", ".join(matched) if matched else "None identified")

            gap_p = doc.add_paragraph()
            gap_p.paragraph_format.space_after = Pt(4)
            r_mg = gap_p.add_run("Skill Gap: ")
            r_mg.font.bold = True
            gap_p.add_run(", ".join(missing) if missing else "None (Full requirements covered)")

            # Projects
            projects = cand.get("projects", [])
            if projects:
                proj_p = doc.add_paragraph()
                proj_p.paragraph_format.space_after = Pt(2)
                r_pj = proj_p.add_run("Projects:")
                r_pj.font.bold = True
                for pr in projects:
                    p_name = pr.get("project_name", "") if isinstance(pr, dict) else str(pr)
                    p_desc = pr.get("description", "") if isinstance(pr, dict) else ""
                    bp = doc.add_paragraph(style='List Bullet')
                    bp.paragraph_format.space_after = Pt(1)
                    bp.add_run(f"{p_name}: {p_desc}")

            # Internships
            internships = cand.get("internships", [])
            if internships:
                int_p = doc.add_paragraph()
                int_p.paragraph_format.space_after = Pt(2)
                r_in = int_p.add_run("Internships / Experience:")
                r_in.font.bold = True
                for it in internships:
                    c_name = it.get("company", "") if isinstance(it, dict) else str(it)
                    c_desc = it.get("description", "") if isinstance(it, dict) else ""
                    bp = doc.add_paragraph(style='List Bullet')
                    bp.paragraph_format.space_after = Pt(1)
                    bp.add_run(f"{c_name}: {c_desc}")

            # Certifications
            certs = cand.get("certifications", [])
            if certs:
                cert_p = doc.add_paragraph()
                cert_p.paragraph_format.space_after = Pt(2)
                r_cr = cert_p.add_run("Certifications:")
                r_cr.font.bold = True
                for ct in certs:
                    ct_name = ct.get("name", "") if isinstance(ct, dict) else str(ct)
                    bp = doc.add_paragraph(style='List Bullet')
                    bp.paragraph_format.space_after = Pt(1)
                    bp.add_run(ct_name)

            # AI Explanation
            explanation = cand.get("ai_explanation", "")
            if explanation:
                exp_p = doc.add_paragraph()
                exp_p.paragraph_format.space_before = Pt(2)
                exp_p.paragraph_format.space_after = Pt(6)
                r_exp_lbl = exp_p.add_run("AI Recommendation Reason: ")
                r_exp_lbl.font.bold = True
                exp_p.add_run(explanation)

            # Separator line
            sep = doc.add_paragraph()
            sep.paragraph_format.space_after = Pt(8)
            r_sep = sep.add_run("─" * 50)
            r_sep.font.color.rgb = RGBColor(226, 232, 240)

    # Disclaimer Footer
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = fp.add_run("TalentMatch • AI-generated recommendation. Final recruitment decisions are made by the recruiter.")
    rf.font.size = Pt(8.5)
    rf.font.color.rgb = RGBColor(148, 163, 184)

    timestamp_str = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    filename = f"TalentMatch_Shortlisted_Candidates_{timestamp_str}_{uuid.uuid4().hex[:6]}.docx"
    output_path = os.path.join(DOCUMENTS_DIR, filename)
    doc.save(output_path)

    return output_path, filename


# ---------------------------------------------------------------------------
# ReportLab PDF Generation
# ---------------------------------------------------------------------------
import io
import html
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfgen import canvas


class NumberedCanvas(canvas.Canvas):
    """
    Two-pass canvas to dynamically compute and draw total page numbers and footer.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_footer(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_footer(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        # Landscape letter width is 792, height is 612
        self.line(28, 28, 792 - 28, 28)
        self.drawString(28, 16, "TalentMatch | AI-Powered Candidate Screening | Confidential Recruitment Document")
        self.drawRightString(792 - 28, 16, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


def create_shortlist_pdf(
    company_name: str,
    college_name: str,
    role_name: str,
    candidates_data: List[Dict[str, Any]]
) -> bytes:
    """
    Generates an executive-ready, beautifully styled PDF document containing ONLY selected shortlisted candidates.
    Uses ReportLab and returns raw PDF bytes.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=28,
        rightMargin=28,
        topMargin=26,
        bottomMargin=38
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "PDFTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E40AF")
    )
    subtitle_style = ParagraphStyle(
        "PDFSubTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B")
    )
    meta_style = ParagraphStyle(
        "PDFMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1E293B")
    )
    th_style = ParagraphStyle(
        "PDFTh",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )
    td_style = ParagraphStyle(
        "PDFTd",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1E293B")
    )
    td_center_style = ParagraphStyle(
        "PDFTdCenter",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1E293B"),
        alignment=1
    )
    td_bold_style = ParagraphStyle(
        "PDFTdBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1D4ED8"),
        alignment=1
    )

    elements = []

    # 1. Header Title & Subtitle
    elements.append(Paragraph("TalentMatch - Shortlisted Candidates", title_style))
    elements.append(Paragraph("AI-POWERED CANDIDATE SCREENING &amp; SHORTLIST REPORT", subtitle_style))
    elements.append(Spacer(1, 8))

    # 2. Metadata Box
    now_date_str = datetime.now().strftime("%B %d, %Y")
    safe_company = html.escape(company_name or "TalentMatch Partner")
    safe_role = html.escape(role_name or "Campus Recruitment")
    safe_college = html.escape(college_name or "Partner Institutions")

    meta_table_data = [
        [
            Paragraph(f"<b>Company Name:</b> {safe_company}", meta_style),
            Paragraph(f"<b>Job Role:</b> {safe_role}", meta_style),
        ],
        [
            Paragraph(f"<b>College Name:</b> {safe_college}", meta_style),
            Paragraph(f"<b>Generation Date:</b> {now_date_str}  |  <b>Selected:</b> {len(candidates_data)} candidate(s)", meta_style),
        ]
    ]

    meta_table = Table(meta_table_data, colWidths=[368, 368])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 10))

    # 3. Candidates Table
    headers = [
        Paragraph("Rank", th_style),
        Paragraph("Candidate Name", th_style),
        Paragraph("Email", th_style),
        Paragraph("Phone", th_style),
        Paragraph("Department", th_style),
        Paragraph("CGPA", th_style),
        Paragraph("Match Score", th_style),
        Paragraph("Matched Skills", th_style),
        Paragraph("Skill Gap", th_style),
        Paragraph("Recruitment Status", th_style),
    ]

    # Column widths sum to 736 pt (landscape letter width 792 - 28*2 margins = 736 pt)
    col_widths = [26, 80, 115, 65, 75, 35, 45, 125, 100, 70]
    table_rows = [headers]

    for idx, cand in enumerate(candidates_data, start=1):
        cand_name = html.escape(str(cand.get("name", "")))
        cand_code = html.escape(str(cand.get("candidate_code", "")))
        cand_email = html.escape(str(cand.get("email", "")))
        cand_phone = html.escape(str(cand.get("phone", "")))
        cand_dept = html.escape(str(cand.get("department", "")))
        
        try:
            cgpa_val = f"{float(cand.get('cgpa', 0.0)):.2f}"
        except (ValueError, TypeError):
            cgpa_val = str(cand.get("cgpa", "-"))

        try:
            score_val = f"{int(round(float(cand.get('match_score', 0.0))))}%"
        except (ValueError, TypeError):
            score_val = str(cand.get("match_score", "-"))

        # Matched skills handling
        m_skills = cand.get("matched_skills", [])
        if isinstance(m_skills, list):
            matched_str = ", ".join(m_skills) if m_skills else "None identified"
        else:
            matched_str = str(m_skills) if m_skills else "None identified"
        safe_matched = html.escape(matched_str)

        # Missing skills handling
        gap_skills = cand.get("missing_skills", [])
        if isinstance(gap_skills, list):
            gap_str = ", ".join(gap_skills) if gap_skills else "None (Full Match)"
        else:
            gap_str = str(gap_skills) if gap_skills else "None (Full Match)"
        safe_gap = html.escape(gap_str)

        status_str = html.escape(str(cand.get("status", "Shortlisted")))

        name_cell_content = f"<b>{cand_name}</b>"
        if cand_code:
            name_cell_content += f"<br/><font color='#64748B' size='6'>{cand_code}</font>"

        row = [
            Paragraph(f"#{idx}", td_center_style),
            Paragraph(name_cell_content, td_style),
            Paragraph(cand_email, td_style),
            Paragraph(cand_phone, td_style),
            Paragraph(cand_dept, td_style),
            Paragraph(cgpa_val, td_center_style),
            Paragraph(score_val, td_bold_style),
            Paragraph(safe_matched, td_style),
            Paragraph(safe_gap, td_style),
            Paragraph(status_str, td_center_style),
        ]
        table_rows.append(row)

    cand_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
    
    table_style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]

    # Alternating row colors
    for r_idx in range(1, len(table_rows)):
        bg = colors.HexColor("#FFFFFF") if r_idx % 2 != 0 else colors.HexColor("#F8FAFC")
        table_style_commands.append(("BACKGROUND", (0, r_idx), (-1, r_idx), bg))

    cand_table.setStyle(TableStyle(table_style_commands))
    elements.append(cand_table)

    doc.build(elements, canvasmaker=NumberedCanvas)
    return buffer.getvalue()
